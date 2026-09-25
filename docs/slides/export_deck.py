"""Render a Beamer PDF into a full-slide-image PowerPoint presentation.

Content editing remains in the .tex source. Requirements: pymupdf, python-pptx.
"""

import argparse
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Inches


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("--outdir", required=True, type=Path)
    args = parser.parse_args()
    args.outdir.mkdir(parents=True, exist_ok=True)
    render_dir = args.outdir / "rendered_slides"
    render_dir.mkdir(exist_ok=True)

    with pymupdf.open(args.pdf) as document:
        if len(document) == 0:
            parser.error("PDF has no pages")
        ratio = document[0].rect.width / document[0].rect.height
        presentation = Presentation()
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = round(presentation.slide_width / ratio)
        presentation.core_properties.title = "心脏超声数据预处理：阶段汇报"
        presentation.core_properties.subject = "LaTeX Beamer PDF 的整页图像放映版"
        presentation.core_properties.author = "eco_prepro"
        for index, page in enumerate(document, start=1):
            if abs(page.rect.width / page.rect.height - ratio) > 0.001:
                parser.error("PDF pages have different aspect ratios")
            png_path = render_dir / f"slide_{index:02d}.png"
            scale = 1920 / page.rect.width
            page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).save(png_path)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(
                str(png_path), 0, 0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
            slide.notes_slide.notes_text_frame.text = (
                "整页图像放映版；修改内容请编辑 LaTeX 源码。\n\n" + page.get_text()
            )
        destination = args.outdir / f"{args.pdf.stem}.pptx"
        presentation.save(destination)
        print(f"Wrote {len(presentation.slides)} slides to {destination}")


if __name__ == "__main__":
    main()
